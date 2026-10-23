from __future__ import annotations

import json
import re
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from html.parser import HTMLParser
from io import BytesIO
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from typing import Any
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PDOK = "https://api.pdok.nl/bzk/locatieserver/search/v3_1"
PARCEL_API = "https://api.pdok.nl/kadaster/brk-kadastrale-kaart/ogc/v1"
VIEWER = "https://document-viewer.dso.kadaster.nl"
PORTAL = "https://omgevingswet.overheid.nl/regels-op-de-kaart"
USER_AGENT = "AdresRegels/0.1 (local research tool)"
UPSTREAM_TIMEOUT = 20
REQUEST_LOCK = threading.Lock()
RATE_LIMIT_LOCK = threading.Lock()
LAST_VIEWER_REQUEST = 0.0

HTML = r"""<!doctype html>
<html lang="nl">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Regels voor uw adres</title>
  <style>
    :root { color-scheme: light; font-family: Arial, sans-serif; color: #162b3b; background: #f2f6f8; }
    body { margin: 0; }
    main { max-width: 1120px; margin: 48px auto; padding: 0 24px; }
    h1 { margin-bottom: 8px; font-size: 2rem; }
    .intro, .notice { color: #465966; line-height: 1.55; }
    .card { margin-top: 24px; padding: 24px; border: 1px solid #d4dfe5; border-radius: 10px; background: white; }
    label { display: block; margin-bottom: 8px; font-weight: 700; }
    .search { display: flex; gap: 12px; }
    input { flex: 1; min-width: 0; padding: 12px; border: 1px solid #8296a3; border-radius: 5px; font-size: 1rem; }
    button { padding: 12px 18px; border: 0; border-radius: 5px; background: #005a8b; color: white; font-size: 1rem; cursor: pointer; }
    button.secondary { background: #267044; }
    button:disabled { cursor: wait; opacity: .6; }
    .export-options { display: flex; align-items: center; gap: 16px; margin-bottom: 12px; }
    .checkbox-label { display: flex; align-items: center; gap: 8px; margin: 0; font-weight: 400; }
    .checkbox-label input { flex: 0 0 auto; width: auto; margin: 0; }
    .document-checkbox { flex: 0 0 auto; width: auto; min-width: auto; margin: 0; padding: 0; border: 0; }
    .version-toggle { display: flex; align-items: center; gap: 8px; margin: 0; font-weight: 400; }
    .version-toggle input {
      appearance: none;
      flex: 0 0 auto;
      width: 38px;
      height: 22px;
      margin: 0;
      padding: 0;
      border: 1px solid #71808a;
      border-radius: 12px;
      background: #cbd3d8;
      cursor: pointer;
      transition: background-color .15s ease;
    }
    .version-toggle input::before {
      display: block;
      width: 16px;
      height: 16px;
      margin: 2px;
      border-radius: 50%;
      background: white;
      content: "";
      transition: transform .15s ease;
    }
    .version-toggle input:checked::before { transform: translateX(16px); }
    .version-toggle input:focus-visible { outline: 3px solid #005a8b; outline-offset: 2px; }
    .version-toggle--draft input:checked { border-color: #d2a900; background: #f3d34a; }
    .version-toggle--adopted input:checked { border-color: #16843b; background: #27a34a; }
    #suggestions { position: relative; z-index: 1; }
    .suggestion { display: block; width: 100%; margin-top: 4px; text-align: left; background: #eef5f8; color: #162b3b; }
    .status { min-height: 1.5em; margin-top: 14px; }
    .error { color: #a32222; }
    .table-wrap { overflow-x: auto; }
    table { width: 100%; border-collapse: collapse; }
    th, td { padding: 12px; border-bottom: 1px solid #d4dfe5; text-align: left; vertical-align: top; }
    th { background: #eef5f8; }
    a { color: #005a8b; }
    .notice { margin-top: 20px; font-size: .9rem; }
    @media (max-width: 640px) {
      main { margin-top: 24px; padding: 0 14px; }
      .card { padding: 16px; }
      .search { flex-direction: column; }
      .export-options { align-items: flex-start; flex-direction: column; }
    }
  </style>
</head>
<body>
  <main>
    <h1>Regels voor uw adres</h1>
    <p class="intro">Zoek een woonadres en bekijk in één overzicht de regelgevende documenten die op die locatie van toepassing zijn. Open de tab ‘Regels’ om de officiële regels te lezen.</p>
    <section class="card" aria-label="Adres zoeken">
      <label for="address">Woonadres</label>
      <div class="search">
        <input id="address" autocomplete="off" placeholder="Bijvoorbeeld: Binnenhof 1, Den Haag" aria-autocomplete="list" aria-controls="suggestions">
        <button id="search" type="button" disabled>Toon documenten</button>
      </div>
      <div id="suggestions" role="listbox"></div>
      <div id="status" class="status" role="status" aria-live="polite"></div>
    </section>
    <section id="results" class="card" hidden>
      <h2 id="results-title">Regels</h2>
      <p id="document-count"></p>
      <div class="export-options">
        <label class="checkbox-label" for="include-chapter-22">
          <input id="include-chapter-22" type="checkbox">
          Hoofdstuk 22 meenemen
        </label>
        <label class="version-toggle version-toggle--draft" for="include-drafts">
          <input id="include-drafts" type="checkbox" role="switch">
          Ontwerp meenemen
        </label>
        <label class="version-toggle version-toggle--adopted" for="include-adopted">
          <input id="include-adopted" type="checkbox" role="switch">
          Vastgesteld meenemen
        </label>
        <button id="download-word" class="secondary" type="button">Download Word met alle artikelen op deze locatie</button>
      </div>
      <div class="table-wrap">
        <table>
          <thead><tr><th>Meenemen in Word</th><th>Document</th><th>Type</th><th>Bevoegd gezag</th><th>In werking sinds</th><th>Bron</th></tr></thead>
          <tbody id="rows"></tbody>
        </table>
      </div>
    </section>
    <p class="notice">De Word-export neemt de volledige tekst op van artikelen die volgens het locatie-filter van Regels op de kaart op het gekozen perceel van toepassing zijn. Gegevens komen van PDOK en het Omgevingsloket. Controleer altijd de officiële bron; dit overzicht geeft geen juridisch advies.</p>
  </main>
  <script>
    const address = document.querySelector("#address");
    const searchButton = document.querySelector("#search");
    const suggestions = document.querySelector("#suggestions");
    const status = document.querySelector("#status");
    const results = document.querySelector("#results");
    const downloadButton = document.querySelector("#download-word");
    const chapter22Checkbox = document.querySelector("#include-chapter-22");
    const includeDraftsCheckbox = document.querySelector("#include-drafts");
    const includeAdoptedCheckbox = document.querySelector("#include-adopted");
    let selected = null;
    let timer;
    let latestQuery = "";

    function setStatus(message, isError = false) {
      status.textContent = message;
      status.classList.toggle("error", isError);
    }

    address.addEventListener("input", () => {
      selected = null;
      searchButton.disabled = true;
      suggestions.replaceChildren();
      results.hidden = true;
      const query = address.value.trim();
      latestQuery = query;
      clearTimeout(timer);
      if (query.length < 3) return;
      timer = setTimeout(async () => {
        try {
          const response = await fetch(`/api/suggest?q=${encodeURIComponent(query)}`);
          const data = await response.json();
          if (!response.ok) throw new Error(data.error || "Adressen ophalen is niet gelukt.");
          if (latestQuery !== query) return;
          for (const item of data.suggestions) {
            const option = document.createElement("button");
            option.type = "button";
            option.className = "suggestion";
            option.role = "option";
            option.textContent = item.label;
            option.addEventListener("click", () => {
              selected = item;
              address.value = item.label;
              suggestions.replaceChildren();
              searchButton.disabled = false;
              setStatus("");
            });
            suggestions.append(option);
          }
          if (!data.suggestions.length) setStatus("Geen adres gevonden. Probeer een straatnaam, huisnummer en plaats.");
        } catch (error) {
          setStatus(error.message, true);
        }
      }, 250);
    });

    searchButton.addEventListener("click", async () => {
      if (!selected) return;
      searchButton.disabled = true;
      results.hidden = true;
      setStatus("Perceel en documenten ophalen…");
      try {
        const response = await fetch("/api/documents", {
          method: "POST",
          headers: { "Content-Type": "application/json" },
          body: JSON.stringify({ id: selected.id, label: selected.label })
        });
        const data = await response.json();
        if (!response.ok) throw new Error(data.error || "De regels ophalen is niet gelukt.");
        document.querySelector("#results-title").textContent = `Documenten voor ${data.address}`;
        const body = document.querySelector("#rows");
        body.replaceChildren();
        for (const doc of data.documents) {
          const row = document.createElement("tr");
          const selectionCell = document.createElement("td");
          const checkbox = document.createElement("input");
          checkbox.type = "checkbox";
          checkbox.className = "document-checkbox";
          checkbox.dataset.documentId = doc.id;
          checkbox.checked = doc.type.trim().toLocaleLowerCase() === "omgevingsplan";
          checkbox.setAttribute("aria-label", `Neem ${doc.title} mee in Word`);
          selectionCell.append(checkbox);
          row.append(selectionCell);
          for (const value of [doc.title, doc.type, doc.authority, doc.inForceSince]) {
            const cell = document.createElement("td");
            cell.textContent = value;
            row.append(cell);
          }
          const sourceCell = document.createElement("td");
          const link = document.createElement("a");
          link.href = doc.url;
          link.target = "_blank";
          link.rel = "noopener noreferrer";
          link.textContent = "Bekijk Regels";
          sourceCell.append(link);
          row.append(sourceCell);
          body.append(row);
        }
        results.hidden = false;
        document.querySelector("#document-count").textContent = `${data.documents.length} regelgevende documenten gevonden.`;
        setStatus(data.documents.length ? "Klaar om alle artikelen met het locatie-filter in Word te zetten." : "Geen regelgevende documenten gevonden voor deze locatie.");
      } catch (error) {
        setStatus(error.message, true);
      } finally {
        searchButton.disabled = false;
      }
    });

    downloadButton.addEventListener("click", async () => {
      if (!selected) return;
      downloadButton.disabled = true;
      setStatus("Alle locatiegefilterde artikelen ophalen en Word-bestand maken. Dit kan even duren…");
      try {
        const response = await fetch("/api/word", {
          method: "POST",
          headers: { "Content-Type": "application/json" },
          body: JSON.stringify({
            id: selected.id,
            label: selected.label,
            includeChapter22: chapter22Checkbox.checked,
            includeDraftVersions: includeDraftsCheckbox.checked,
            includeAdoptedVersions: includeAdoptedCheckbox.checked,
            selectedDocumentIds: Array.from(
              document.querySelectorAll("#rows .document-checkbox:checked"),
              checkbox => checkbox.dataset.documentId
            )
          })
        });
        if (!response.ok) {
          const data = await response.json();
          throw new Error(data.error || "Het Word-bestand kon niet worden gemaakt.");
        }
        const blob = await response.blob();
        const url = URL.createObjectURL(blob);
        const link = document.createElement("a");
        link.href = url;
        const encodedFilename = response.headers.get("X-Download-Filename");
        link.download = encodedFilename
          ? decodeURIComponent(encodedFilename)
          : "RPL_adres.docx";
        link.click();
        URL.revokeObjectURL(url);
        setStatus("Het Word-bestand is gedownload.");
      } catch (error) {
        setStatus(error.message, true);
      } finally {
        downloadButton.disabled = false;
      }
    });
  </script>
</body>
</html>
"""


class UpstreamError(Exception):
    pass


def request_json(url: str, method: str = "GET", body: dict[str, Any] | None = None) -> dict[str, Any]:
    global LAST_VIEWER_REQUEST
    parsed = urllib.parse.urlsplit(url)
    if parsed.hostname not in {"api.pdok.nl", "document-viewer.dso.kadaster.nl"}:
        raise UpstreamError("Een externe API gaf een ongeldige verwijzing terug.")
    if parsed.hostname == "document-viewer.dso.kadaster.nl":
        with RATE_LIMIT_LOCK:
            delay = 0.3 - (time.monotonic() - LAST_VIEWER_REQUEST)
            if delay > 0:
                time.sleep(delay)
            LAST_VIEWER_REQUEST = time.monotonic()
    data = json.dumps(body).encode("utf-8") if body is not None else None
    is_ihr = parsed.path.startswith("/bff/ihr/")
    accept = (
        "application/geo+json, application/json, */*"
        if parsed.hostname == "api.pdok.nl"
        else "*/*" if is_ihr else "application/json, application/hal+json, */*"
    )
    headers = {
        "Accept": accept,
        "User-Agent": USER_AGENT,
        "Referer": VIEWER + "/",
        "Origin": VIEWER,
    }
    if data is not None:
        headers["Content-Type"] = "application/json"
        if parsed.hostname == "document-viewer.dso.kadaster.nl" and (
            "geometrie" in body or "geometry" in body or "_geo" in body
        ):
            headers["Content-Crs"] = (
                "epsg:28992" if is_ihr else "http://www.opengis.net/def/crs/EPSG/0/28992"
            )
    request = urllib.request.Request(url, data=data, headers=headers, method=method)
    for attempt in range(3):
        try:
            with urllib.request.urlopen(request, timeout=UPSTREAM_TIMEOUT) as response:
                payload = response.read()
            break
        except urllib.error.HTTPError as error:
            detail = error.read(1000).decode("utf-8", errors="replace").strip()
            if error.code in {429, 503} and attempt < 2:
                retry_after = error.headers.get("Retry-After", "")
                delay = int(retry_after) if retry_after.isdigit() else 2 ** (attempt + 1)
                time.sleep(delay)
                continue
            suffix = f": {detail}" if detail else ""
            raise UpstreamError(
                f"{parsed.hostname}{parsed.path} gaf HTTP {error.code} terug{suffix}"
            ) from error
        except urllib.error.URLError as error:
            raise UpstreamError(
                f"{parsed.hostname}{parsed.path} is niet bereikbaar: {error.reason}"
            ) from error
    try:
        result = json.loads(payload)
    except (json.JSONDecodeError, UnicodeDecodeError) as error:
        raise UpstreamError("Een gegevensbron gaf geen geldige JSON terug.") from error
    if not isinstance(result, dict):
        raise UpstreamError("Een gegevensbron gaf een onverwacht antwoord terug.")
    return result


def get_suggestions(query: str) -> list[dict[str, str]]:
    params = urllib.parse.urlencode(
        {"fq": "type:adres", "q": query, "rows": "8", "fl": "id,type,weergavenaam"}
    )
    response = request_json(f"{PDOK}/suggest?{params}")
    docs = response.get("response", {}).get("docs", [])
    return [
        {"id": item["id"], "label": item["weergavenaam"]}
        for item in docs
        if item.get("type") == "adres" and item.get("id") and item.get("weergavenaam")
    ]


def parcel_geometry(address: dict[str, Any]) -> dict[str, Any]:
    parcels = address.get("gekoppeld_perceel") or []
    parsed: list[tuple[str, str, int]] = []
    for parcel in parcels:
        match = re.fullmatch(r"([A-Z0-9]+)-([A-Z])-([0-9]+)", parcel)
        if match:
            parsed.append((match.group(1), match.group(2), int(match.group(3))))
    if not parsed:
        raise UpstreamError("PDOK heeft geen kadastraal perceel aan dit adres gekoppeld.")

    clauses = [
        "(akr_kadastrale_gemeente_code_waarde = "
        + cql_string(code)
        + " AND sectie = "
        + cql_string(section)
        + f" AND perceelnummer = {number})"
        for code, section, number in parsed
    ]
    params = urllib.parse.urlencode(
        {
            "f": "json",
            "limit": str(max(10, len(parsed) * 2)),
            "crs": "http://www.opengis.net/def/crs/EPSG/0/28992",
            "filter-lang": "cql2-text",
            "filter": " OR ".join(clauses),
        }
    )
    feature_collection = request_json(f"{PARCEL_API}/collections/perceel/items?{params}")
    geometries: list[dict[str, Any]] = []
    for feature in feature_collection.get("features", []):
        if not isinstance(feature, dict) or not isinstance(feature.get("geometry"), dict):
            continue
        parcel_geometry = feature["geometry"]
        if parcel_geometry.get("type") == "MultiPolygon":
            geometries.extend(
                {"type": "Polygon", "coordinates": polygon}
                for polygon in parcel_geometry.get("coordinates", [])
            )
        elif parcel_geometry.get("type") == "Polygon":
            geometries.append(parcel_geometry)
    if not geometries:
        raise UpstreamError("PDOK kon de gekoppelde perceelgrens niet ophalen.")

    centroid = address.get("centroide_rd", "")
    point_match = re.fullmatch(r"POINT\(\s*(-?\d+(?:\.\d+)?)\s+(-?\d+(?:\.\d+)?)\s*\)", centroid)
    if point_match:
        geometries.append(
            {"type": "Point", "coordinates": [float(point_match.group(1)), float(point_match.group(2))]}
        )
    return {"type": "GeometryCollection", "geometries": geometries}


def cql_string(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def display_value(value: Any) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        labels = [display_value(item) for item in value]
        return ", ".join(label for label in labels if label)
    if isinstance(value, dict):
        for key in ("naam", "weergavenaam", "label", "officieleNaam"):
            if value.get(key):
                return display_value(value[key])
    return ""


def get_location_documents(
    address_id: str, label: str
) -> tuple[str, dict[str, Any], list[dict[str, Any]]]:
    if not re.fullmatch(r"adr-[a-f0-9]+", address_id):
        raise UpstreamError("De gekozen adresverwijzing is ongeldig.")
    lookup = request_json(f"{PDOK}/lookup?{urllib.parse.urlencode({'id': address_id})}")
    docs = lookup.get("response", {}).get("docs", [])
    if not docs or docs[0].get("type") != "adres":
        raise UpstreamError("Het gekozen adres kon niet door PDOK worden bevestigd.")
    address = docs[0]
    geometry = parcel_geometry(address)
    search_url = (
        f"{VIEWER}/bff/ois/ontsluiten/v2/documenten/_zoek?"
        + urllib.parse.urlencode(
            {
                "page": "0",
                "size": "200",
                "_sort": "DESC",
                "synchroniseerMetTileset": "Actueel",
            }
        )
    )
    layers = ("GEMEENTE", "PROVINCIE", "WATERSCHAP", "RIJK")

    def search_layer(layer: str) -> list[dict[str, Any]]:
        data = request_json(
            search_url,
            method="POST",
            body={
                "regelgevingOfOverig": "REGELGEVING",
                "geometrie": geometry,
                "bestuurslaag": layer,
            },
        )
        return data.get("_embedded", {}).get("documenten", [])

    with ThreadPoolExecutor(max_workers=4) as pool:
        documents = [doc for layer_docs in pool.map(search_layer, layers) for doc in layer_docs]
    unique_documents = {
        document["identificatie"]: document
        for document in documents
        if isinstance(document, dict) and isinstance(document.get("identificatie"), str)
    }
    return address.get("weergavenaam", label), geometry, list(unique_documents.values())


def fetch_documents(address_id: str, label: str) -> dict[str, Any]:
    address_label, _, documents = get_location_documents(address_id, label)
    results: list[dict[str, str]] = []
    for document in documents:
        identifier = document["identificatie"]
        slug = "_" + identifier.strip("/").replace("/", "_")
        if not re.fullmatch(r"_[A-Za-z0-9_.-]+", slug):
            continue
        location_query = urllib.parse.urlencode({"locatie": address_label})
        title = display_value(document.get("titel")) or "Onbekend document"
        authority = display_value(document.get("aangeleverdDoorEen"))
        in_force = document.get("inwerkingVanaf") or document.get("geldigVanaf") or ""
        results.append(
            {
                "id": identifier,
                "title": title,
                "type": display_value(document.get("type")) or "Onbekend",
                "authority": authority or "Onbekend",
                "inForceSince": str(in_force)[:10] if in_force else "Onbekend",
                "url": f"{PORTAL}/documenten/{slug}/regels?{location_query}",
            }
        )
    results.sort(key=lambda item: (item["authority"], item["title"]))
    return {"address": address_label, "documents": results}


class MarkupTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"br", "p", "li", "div", "al", "lijst", "table", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"p", "li", "div", "al", "lijst", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        self.parts.append(data)

    def text(self) -> str:
        lines = [" ".join(part.split()) for part in "".join(self.parts).splitlines()]
        return "\n".join(line for line in lines if line)


class MarkupBlockParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[dict[str, Any]] = []
        self.text_parts: list[str] = []
        self.table_rows: list[list[str]] | None = None
        self.table_depth = 0
        self.current_row: list[str] = []
        self.current_cell: list[str] | None = None

    def flush_paragraph(self) -> None:
        text = " ".join("".join(self.text_parts).split())
        if text:
            self.blocks.append({"type": "paragraph", "text": text})
        self.text_parts.clear()

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag == "table":
            if self.table_depth == 0:
                self.flush_paragraph()
                self.table_rows = []
            self.table_depth += 1
        elif self.table_depth:
            if tag == "tr":
                self.current_row = []
            elif tag in {"td", "th"}:
                self.current_cell = []
            elif tag == "br" and self.current_cell is not None:
                self.current_cell.append(" ")
        elif tag in {"br", "p", "li", "div", "al", "lijst"}:
            self.flush_paragraph()

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if self.table_depth:
            if tag in {"td", "th"} and self.current_cell is not None:
                self.current_row.append(" ".join("".join(self.current_cell).split()))
                self.current_cell = None
            elif tag == "tr" and self.table_rows is not None:
                self.table_rows.append(self.current_row)
                self.current_row = []
            elif tag == "table":
                self.table_depth -= 1
                if self.table_depth == 0 and self.table_rows is not None:
                    rows = self.table_rows
                    self.table_rows = None
                    if rows:
                        self.blocks.append({"type": "table", "rows": rows})
        elif tag in {"p", "li", "div", "al", "lijst"}:
            self.flush_paragraph()

    def handle_data(self, data: str) -> None:
        if self.current_cell is not None:
            self.current_cell.append(data)
        elif not self.table_depth:
            self.text_parts.append(data)

    def parse(self, value: str) -> list[dict[str, Any]]:
        self.feed(value)
        self.close()
        self.flush_paragraph()
        return self.blocks


def markup_text(value: Any) -> str:
    if not isinstance(value, str) or not value:
        return ""
    parser = MarkupTextParser()
    parser.feed(value)
    parser.close()
    return parser.text()


def markup_blocks(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, str) or not value:
        return []
    return MarkupBlockParser().parse(value)


def content_blocks_text(blocks: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    for block in blocks:
        if block.get("type") == "paragraph":
            text = block.get("text")
            if isinstance(text, str) and text:
                parts.append(text)
        elif block.get("type") == "table":
            rows = block.get("rows", [])
            if isinstance(rows, list):
                parts.extend(
                    " | ".join(str(cell) for cell in row)
                    for row in rows
                    if isinstance(row, list)
                )
    return "\n".join(parts)


def add_query_parameter(url: str, name: str, value: str) -> str:
    parsed = urllib.parse.urlsplit(url)
    query = urllib.parse.parse_qsl(parsed.query, keep_blank_values=True)
    if not any(key == name for key, _ in query):
        query.append((name, value))
    return urllib.parse.urlunsplit(
        (parsed.scheme, parsed.netloc, parsed.path, urllib.parse.urlencode(query), parsed.fragment)
    )


def regulation_slug(identifier: str) -> str:
    slug = "_" + identifier.strip("/").replace("/", "_")
    if not re.fullmatch(r"_[A-Za-z0-9_.-]+", slug):
        raise UpstreamError("Een document heeft een ongeldige identificatie.")
    return slug


def annotation_location_data(
    annotation_url: str, geometry: dict[str, Any]
) -> dict[str, Any]:
    parsed = urllib.parse.urlsplit(annotation_url)
    if not (
        parsed.path.endswith("/regeltekstannotaties")
        or parsed.path.endswith("/regeltekstannotaties/renvooi")
    ):
        raise UpstreamError("De officiële bron gaf geen link naar de regelannotaties.")
    location_url = urllib.parse.urlunsplit(
        (parsed.scheme, parsed.netloc, parsed.path + "/_zoek", parsed.query, parsed.fragment)
    )
    return request_json(
        f"{VIEWER}/bff/logic/annotaties-op-locatie/_zoek",
        method="POST",
        body={"url": location_url, "geometry": geometry},
    )


def document_component_children(component: dict[str, Any]) -> list[Any]:
    embedded = component.get("_embedded", {})
    if not isinstance(embedded, dict):
        return []
    children: list[Any] = []
    for key in ("documentComponenten", "ontwerpDocumentComponenten"):
        values = embedded.get(key, [])
        if isinstance(values, list):
            children.extend(values)
    return children


def selected_ozon_rule_texts(
    regulation: dict[str, Any], geometry: dict[str, Any]
) -> tuple[dict[str, Any], set[str]]:
    links = regulation.get("_links", {})
    annotation_link = links.get("annotaties", {}).get("href")
    structure_link = links.get("documentstructuur", {}).get("href")
    if not isinstance(annotation_link, str) or not isinstance(structure_link, str):
        raise UpstreamError("Het document bevat geen link naar de regels en documentstructuur.")

    annotation_url = add_query_parameter(annotation_link, "locatieSelectie", "primair")
    annotations = request_json(annotation_url)
    location = annotation_location_data(annotation_url, geometry)
    selected_refs = set(location.get("geselecteerdeLocatieRefs", []))
    selected_activity_locations = {
        identifier
        for activity in location.get("activiteiten", [])
        if isinstance(activity, dict)
        for identifier in activity.get("activiteitLocatieaanduidingen", [])
        if isinstance(identifier, str)
    }

    relevant_text_refs: set[str] = set()
    for rule in annotations.get("regelsVoorIedereen", []):
        if not isinstance(rule, dict):
            continue
        location_matches = bool(set(rule.get("locatieRefs", [])).intersection(selected_refs))
        activity_matches = any(
            isinstance(item, dict)
            and item.get("identificatie") in selected_activity_locations
            for item in rule.get("activiteitLocatieaanduidingen", [])
        )
        if (location_matches or activity_matches) and isinstance(rule.get("regeltekstRef"), str):
            relevant_text_refs.add(rule["regeltekstRef"])

    text_locations = {
        item.get("identificatie"): item.get("wId")
        for item in annotations.get("regelteksten", [])
        if isinstance(item, dict)
        and item.get("identificatie") in relevant_text_refs
        and isinstance(item.get("wId"), str)
    }
    target_component_ids = set(text_locations.values())
    structure = request_json(structure_link)
    articles: dict[str, dict[str, Any]] = {}

    def visit(component: dict[str, Any], current_article: dict[str, Any] | None = None) -> None:
        if component.get("type") == "ARTIKEL":
            current_article = component
        if current_article and component.get("identificatie") in target_component_ids:
            articles.setdefault(current_article["identificatie"], current_article)
        for child in document_component_children(component):
            if isinstance(child, dict):
                visit(child, current_article)

    for root in document_component_children(structure):
        if isinstance(root, dict):
            visit(root)
    return structure, set(articles)


def get_structure_articles(
    structure: dict[str, Any],
    selected_article_ids: set[str],
    document_title: str,
    authority: str,
    source_url: str,
    only_changed_components: bool = False,
) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    slug = urllib.parse.urlsplit(source_url).path.rsplit("/", 2)[-2]

    def collect_content(
        component: dict[str, Any],
        blocks: list[dict[str, Any]],
        seen_blocks: set[str],
    ) -> None:
        for key in ("kop", "inhoud"):
            for block in markup_blocks(component.get(key)):
                block_key = json.dumps(block, ensure_ascii=False, sort_keys=True)
                if block_key not in seen_blocks:
                    seen_blocks.add(block_key)
                    blocks.append(block)
        for child in document_component_children(component):
            if isinstance(child, dict):
                collect_content(child, blocks, seen_blocks)

    section_types = {"HOOFDSTUK", "AFDELING", "PARAGRAAF"}

    def has_version_change(component: dict[str, Any]) -> bool:
        if component.get("bevatRenvooi") or component.get("bevatOntwerpInformatie"):
            return True
        return any(
            has_version_change(child)
            for child in document_component_children(component)
            if isinstance(child, dict)
        )

    def structure_heading(component: dict[str, Any]) -> str:
        heading = markup_text(component.get("kop")) or display_value(component.get("titel"))
        heading = " ".join(heading.split())
        heading = re.sub(
            r"(?i)\b(hoofdstuk|afdeling|paragraaf)\s*(\d+(?:\.\d+)*)",
            r"\1 \2",
            heading,
        )
        return re.sub(r"(?<=\d)(?=[^\W\d_])", " ", heading)

    def visit(
        component: dict[str, Any],
        hierarchy: list[dict[str, str]],
    ) -> None:
        component_type = component.get("type")
        current_hierarchy = hierarchy
        if component_type in section_types:
            title = structure_heading(component)
            if title:
                current_hierarchy = [
                    *hierarchy,
                    {
                        "type": component_type,
                        "title": title,
                        "identifier": str(component.get("identificatie", "")),
                    },
                ]
        if (
            component.get("identificatie") in selected_article_ids
            and component.get("type") == "ARTIKEL"
            and (not only_changed_components or has_version_change(component))
        ):
            blocks: list[dict[str, Any]] = []
            collect_content(component, blocks, set())
            number = markup_text(component.get("kop"))
            heading = f"Artikel {number}".strip() if number else "Artikel"
            if blocks and blocks[0].get("type") == "paragraph" and blocks[0].get("text") == number:
                blocks.pop(0)
            results.append(
                {
                    "key": component["identificatie"],
                    "document": document_title,
                    "authority": authority,
                    "heading": heading,
                    "text": content_blocks_text(blocks),
                    "blocks": blocks,
                    "hierarchy": current_hierarchy,
                    "url": (
                        f"{source_url}#{urllib.parse.quote(slug + '-' + component['identificatie'], safe='-_.')}"
                    ),
                }
            )
        for child in document_component_children(component):
            if isinstance(child, dict):
                visit(child, current_hierarchy)

    for root in document_component_children(structure):
        if isinstance(root, dict):
            visit(root, [])
    return results


def fetch_ozon_regulation_articles(
    regulation: dict[str, Any],
    geometry: dict[str, Any],
    address_label: str,
    seen: set[str],
    depth: int = 0,
    version_label: str | None = None,
    source_slug: str | None = None,
) -> list[dict[str, Any]]:
    identifier = regulation.get("identificatie")
    if not isinstance(identifier, str):
        raise UpstreamError("Een omgevingsdocument heeft geen identificatie.")
    slug = source_slug or regulation_slug(identifier)
    if slug in seen:
        return []
    if depth > 3:
        raise UpstreamError("De verwijzingen naar tijdelijke documentdelen zijn te diep genest.")
    seen.add(slug)

    title = (
        display_value(regulation.get("officieleTitel"))
        or display_value(regulation.get("citeerTitel"))
        or display_value(regulation.get("titel"))
        or identifier
    )
    authority = display_value(regulation.get("aangeleverdDoorEen")) or "Onbekend bevoegd gezag"
    structure, article_ids = selected_ozon_rule_texts(regulation, geometry)
    query_params: list[tuple[str, str]] = [("locatie", address_label)]
    if version_label == "Vastgesteld":
        self_href = regulation.get("_links", {}).get("self", {}).get("href")
        if isinstance(self_href, str):
            self_query = urllib.parse.parse_qs(
                urllib.parse.urlsplit(self_href).query,
                keep_blank_values=True,
            )
            for name in ("geldigOp", "inWerkingOp", "beschikbaarOp"):
                values = self_query.get(name, [])
                if values:
                    query_params.append((name, values[0]))
    location_query = urllib.parse.urlencode(query_params)
    source_url = f"{PORTAL}/documenten/{slug}/regels?{location_query}"
    articles = get_structure_articles(
        structure,
        article_ids,
        title,
        authority,
        source_url,
        only_changed_components=version_label == "Ontwerp",
    )
    if version_label:
        for article in articles:
            article["version_label"] = version_label

    for temporary_part in regulation.get("_links", {}).get("tijdelijkDelen", []):
        href = temporary_part.get("href") if isinstance(temporary_part, dict) else None
        if not isinstance(href, str):
            continue
        part = request_json(href)
        articles.extend(
            fetch_ozon_regulation_articles(
                part,
                geometry,
                address_label,
                seen,
                depth + 1,
                version_label,
            )
        )
    return articles


def changed_version_articles(
    current_articles: list[dict[str, Any]],
    version_articles: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    current_by_key = {
        article["key"]: article
        for article in current_articles
        if isinstance(article.get("key"), str)
    }
    current_by_heading = {
        (
            tuple(item.get("title", "") for item in article.get("hierarchy", [])),
            article.get("heading", ""),
        ): article
        for article in current_articles
    }
    changed: list[dict[str, Any]] = []
    for article in version_articles:
        current = current_by_key.get(article.get("key"))
        if current is None:
            current = current_by_heading.get(
                (
                    tuple(item.get("title", "") for item in article.get("hierarchy", [])),
                    article.get("heading", ""),
                )
            )
        if current is None or current.get("text") != article.get("text"):
            changed.append(article)
    return changed


def collect_ihr_articles(value: Any, target: dict[str, str]) -> None:
    if isinstance(value, dict):
        article_list = value.get("artikelen")
        if isinstance(article_list, list):
            for article in article_list:
                if (
                    isinstance(article, dict)
                    and isinstance(article.get("id"), str)
                    and isinstance(article.get("titel"), str)
                ):
                    target.setdefault(article["id"], article["titel"])
        for child in value.values():
            collect_ihr_articles(child, target)
    elif isinstance(value, list):
        for child in value:
            collect_ihr_articles(child, target)


def ihr_text_pages(plan_id: str) -> dict[str, dict[str, Any]]:
    encoded_id = urllib.parse.quote(plan_id, safe=".-_")
    next_url = f"{VIEWER}/bff/ihr/v4/plannen/{encoded_id}/teksten?pageSize=100"
    texts: dict[str, dict[str, Any]] = {}
    visited: set[str] = set()
    while next_url:
        if next_url in visited or len(visited) >= 100:
            raise UpstreamError(f"De tekstpagina's van plan {plan_id} konden niet volledig worden geladen.")
        visited.add(next_url)
        page = request_json(next_url)
        for item in page.get("_embedded", {}).get("teksten", []):
            if isinstance(item, dict) and isinstance(item.get("id"), str):
                texts[item["id"]] = item
        next_link = page.get("_links", {}).get("next", {}).get("href")
        next_url = next_link if isinstance(next_link, str) else ""
    return texts


def ihr_article_blocks(
    article_id: str, texts: dict[str, dict[str, Any]]
) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    visited: set[str] = set()

    def visit(text_id: str) -> None:
        if text_id in visited:
            return
        visited.add(text_id)
        item = texts.get(text_id)
        if item is None:
            return
        title = item.get("titel")
        if text_id != article_id and isinstance(title, str) and title:
            blocks.append({"type": "paragraph", "text": title})
        blocks.extend(markup_blocks(item.get("inhoud")))
        for child in item.get("_links", {}).get("children", []):
            href = child.get("href") if isinstance(child, dict) else None
            if isinstance(href, str):
                child_id = urllib.parse.unquote(urllib.parse.urlsplit(href).path.rsplit("/", 1)[-1])
                visit(child_id)

    visit(article_id)
    return blocks


def ihr_article_text(article_id: str, texts: dict[str, dict[str, Any]]) -> str:
    return content_blocks_text(ihr_article_blocks(article_id, texts))


def ihr_article_hierarchy(article: dict[str, Any]) -> list[dict[str, str]]:
    hierarchy: list[dict[str, str]] = []
    for item in article.get("kruimelpad", []):
        if not isinstance(item, dict):
            continue
        title = item.get("titel")
        if not isinstance(title, str):
            continue
        match = re.match(r"^(Hoofdstuk|Afdeling|Paragraaf)\b", title, re.IGNORECASE)
        if match:
            hierarchy.append(
                {
                    "type": match.group(1).upper(),
                    "title": title,
                    "identifier": str(item.get("id", "")),
                }
            )
    return hierarchy


def fetch_ihr_plan_articles(
    document: dict[str, Any], geometry: dict[str, Any], address_label: str
) -> list[dict[str, Any]]:
    plan_id = document["identificatie"]
    if not re.fullmatch(r"NL\.IMRO\.[A-Za-z0-9._-]+", plan_id):
        raise UpstreamError("Het legacyplan heeft geen herkenbare IMRO-identificatie.")
    encoded_id = urllib.parse.quote(plan_id, safe=".-_")
    planobjects = request_json(
        f"{VIEWER}/bff/ihr/v4/plannen/{encoded_id}/planobjecten/_zoek",
        method="POST",
        body={"_geo": {"intersectAndNotTouches": geometry}},
    )
    selected_articles: dict[str, str] = {}
    collect_ihr_articles(planobjects.get("_embedded", {}), selected_articles)
    if not selected_articles:
        return []
    texts = ihr_text_pages(plan_id)
    slug = regulation_slug(plan_id)
    location_query = urllib.parse.urlencode({"locatie": address_label})
    source_url = f"{PORTAL}/documenten/{slug}/regels?{location_query}"
    authority = display_value(document.get("aangeleverdDoorEen")) or "Onbekend bevoegd gezag"
    title = display_value(document.get("titel")) or "Onbekend document"
    articles: list[dict[str, Any]] = []
    for article_id, article_title in selected_articles.items():
        text_item = texts.get(article_id, {})
        blocks = ihr_article_blocks(article_id, texts)
        text = content_blocks_text(blocks)
        if not text:
            raise UpstreamError(f"De tekst van “{article_title}” in {title} ontbreekt in de bron.")
        articles.append(
            {
                "key": article_id,
                "document": title,
                "authority": authority,
                "heading": article_title,
                "text": text,
                "blocks": blocks,
                "hierarchy": ihr_article_hierarchy(text_item),
                "url": source_url,
            }
        )
    return articles


def applicable_articles(
    address_id: str,
    label: str,
    selected_document_ids: set[str] | None = None,
    include_draft_versions: bool = False,
    include_adopted_versions: bool = False,
) -> tuple[str, list[dict[str, Any]], list[str]]:
    address_label, geometry, documents = get_location_documents(address_id, label)
    if selected_document_ids is None:
        selected_document_ids = {
            document["identificatie"]
            for document in documents
            if display_value(document.get("type")).casefold() == "omgevingsplan"
        }
    available_document_ids = {
        document["identificatie"]
        for document in documents
        if isinstance(document.get("identificatie"), str)
    }
    unknown_document_ids = selected_document_ids - available_document_ids
    if unknown_document_ids:
        raise UpstreamError("Een geselecteerd document hoort niet (meer) bij dit adres.")
    all_articles: list[dict[str, Any]] = []
    failures: list[str] = []
    seen_regulations: set[str] = set()

    for document in documents:
        identifier = document["identificatie"]
        if identifier not in selected_document_ids:
            continue
        title = display_value(document.get("titel")) or identifier
        try:
            if identifier.startswith("/akn/"):
                slug = regulation_slug(identifier)
                regulation = request_json(
                    f"{VIEWER}/bff/ozon/presenteren/v8/regelingen/{slug}"
                )
                current_articles = fetch_ozon_regulation_articles(
                    regulation, geometry, address_label, seen_regulations
                )
                all_articles.extend(current_articles)
                metadata = document.get("omgevingsdocumentMetadata", {})
                if not isinstance(metadata, dict):
                    metadata = {}

                if include_draft_versions:
                    drafts = metadata.get("gerelateerdeOntwerpen", [])
                    for draft in drafts if isinstance(drafts, list) else []:
                        if not isinstance(draft, dict):
                            continue
                        version_slug = draft.get("uriIdentificatie")
                        if not isinstance(version_slug, str) or not re.fullmatch(
                            r"_[A-Za-z0-9_.-]+", version_slug
                        ):
                            failures.append(
                                f"{title} (ontwerp): de versie heeft geen herkenbare identificatie."
                            )
                            continue
                        try:
                            draft_url = (
                                f"{VIEWER}/bff/ozon/presenteren/v8/ontwerpregelingen/"
                                f"{urllib.parse.quote(version_slug, safe='_.-')}"
                            )
                            draft_regulation = request_json(draft_url)
                            all_articles.extend(
                                fetch_ozon_regulation_articles(
                                    draft_regulation,
                                    geometry,
                                    address_label,
                                    set(),
                                    version_label="Ontwerp",
                                    source_slug=version_slug,
                                )
                            )
                        except UpstreamError as error:
                            draft_title = (
                                display_value(draft.get("besluitCiteertitel"))
                                or display_value(draft.get("titel"))
                                or "ontwerpbesluit"
                            )
                            failures.append(f"{draft_title}: {error}")

                if include_adopted_versions:
                    adopted_versions = metadata.get("gerelateerdeBesluitVersies", [])
                    for version in adopted_versions if isinstance(adopted_versions, list) else []:
                        if not isinstance(version, dict) or version.get("isToekomstig") is not True:
                            continue
                        effective_date = version.get("inwerkingVanaf")
                        valid_date = version.get("geldigVanaf") or effective_date
                        if (
                            not isinstance(effective_date, str)
                            or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", effective_date)
                            or not isinstance(valid_date, str)
                            or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", valid_date)
                        ):
                            failures.append(
                                f"{title} (vastgesteld): de toekomstige versie heeft geen geldige ingangsdatum."
                            )
                            continue
                        try:
                            future_url = (
                                f"{VIEWER}/bff/ozon/presenteren/v8/regelingen/{slug}?"
                                + urllib.parse.urlencode(
                                    {
                                        "geldigOp": valid_date,
                                        "inWerkingOp": effective_date,
                                    }
                                )
                            )
                            future_regulation = request_json(future_url)
                            future_articles = fetch_ozon_regulation_articles(
                                future_regulation,
                                geometry,
                                address_label,
                                set(),
                                version_label="Vastgesteld",
                            )
                            all_articles.extend(
                                changed_version_articles(current_articles, future_articles)
                            )
                        except UpstreamError as error:
                            version_title = (
                                display_value(version.get("titel"))
                                or display_value(version.get("officieleTitel"))
                                or "vastgesteld besluit"
                            )
                            failures.append(f"{version_title}: {error}")
            elif identifier.startswith("NL.IMRO."):
                all_articles.extend(fetch_ihr_plan_articles(document, geometry, address_label))
            else:
                failures.append(f"{title}: documenttype niet ondersteund ({identifier}).")
        except UpstreamError as error:
            failures.append(f"{title}: {error}")

    unique: dict[tuple[str, str, str], dict[str, Any]] = {}
    for article in all_articles:
        unique[
            (
                article["document"],
                article.get("version_label", ""),
                article.get("key", article["heading"]),
            )
        ] = article
    return address_label, list(unique.values()), failures


def add_hyperlink(paragraph: Any, text: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def is_chapter_22(article: dict[str, Any]) -> bool:
    return any(
        item.get("type") == "HOOFDSTUK"
        and re.search(r"(?<!\d)22(?!\d)", f"{item.get('title', '')} {item.get('identifier', '')}")
        for item in article.get("hierarchy", [])
    )


def create_word_document(
    address_id: str,
    label: str,
    include_chapter_22: bool = False,
    selected_document_ids: set[str] | None = None,
    include_draft_versions: bool = False,
    include_adopted_versions: bool = False,
) -> tuple[bytes, int, int]:
    address_label, articles, failures = applicable_articles(
        address_id,
        label,
        selected_document_ids,
        include_draft_versions,
        include_adopted_versions,
    )
    if not include_chapter_22:
        articles = [article for article in articles if not is_chapter_22(article)]
    word = Document()
    section = word.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    normal = word.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    word.add_heading(f"Regels voor {address_label}", 0)
    word.add_paragraph(
        f"Locatiegefilterde artikelen uit Regels op de kaart. Gemaakt op "
        f"{datetime.now().astimezone().strftime('%d-%m-%Y %H:%M')}."
    )
    word.add_paragraph(
        "De export bevat de volledige tekst van artikelen waarvan een of meer regelonderdelen "
        "volgens de officiële locatiegegevens op het gekoppelde perceel van toepassing zijn. "
        "Controleer de bronlink voor actuele versies, kaarten en eventuele uitzonderingen."
    )
    word.add_paragraph(f"Aantal gevonden artikelen: {len(articles)}")

    current_document = ""
    current_hierarchy: list[dict[str, str]] = []
    for article in articles:
        if article["document"] != current_document:
            current_document = article["document"]
            word.add_heading(current_document, level=1)
            word.add_paragraph(f"Bevoegd gezag: {article['authority']}")
            current_hierarchy = []
        hierarchy = article.get("hierarchy", [])
        shared_levels = 0
        while (
            shared_levels < len(current_hierarchy)
            and shared_levels < len(hierarchy)
            and current_hierarchy[shared_levels] == hierarchy[shared_levels]
        ):
            shared_levels += 1
        for depth, item in enumerate(hierarchy[shared_levels:], start=shared_levels):
            word.add_heading(item["title"], level=min(depth + 2, 9))
        current_hierarchy = hierarchy
        version_label = article.get("version_label")
        if version_label in {"Ontwerp", "Vastgesteld"}:
            marker = word.add_paragraph()
            marker_run = marker.add_run(version_label)
            marker_run.bold = True
            marker_run.font.highlight_color = (
                WD_COLOR_INDEX.YELLOW
                if version_label == "Ontwerp"
                else WD_COLOR_INDEX.BRIGHT_GREEN
            )
        word.add_heading(article["heading"], level=min(len(hierarchy) + 2, 9))
        blocks = article.get("blocks", [])
        if blocks:
            for block in blocks:
                if block.get("type") == "paragraph":
                    text = block.get("text")
                    if isinstance(text, str) and text.strip():
                        word.add_paragraph(text.strip())
                elif block.get("type") == "table":
                    rows = block.get("rows", [])
                    column_count = max(
                        (len(row) for row in rows if isinstance(row, list)),
                        default=0,
                    )
                    if column_count:
                        table = word.add_table(rows=0, cols=column_count)
                        table.style = "Table Grid"
                        for row in rows:
                            if not isinstance(row, list):
                                continue
                            cells = table.add_row().cells
                            for index, value in enumerate(row):
                                cells[index].text = str(value)
        else:
            for paragraph_text in article["text"].splitlines():
                if paragraph_text.strip():
                    word.add_paragraph(paragraph_text.strip())
        source = word.add_paragraph()
        source.add_run("Officiële bron: ")
        add_hyperlink(source, "Bekijk Regels op de kaart", article["url"])

    if failures:
        word.add_heading("Documenten die niet volledig konden worden verwerkt", level=1)
        word.add_paragraph(
            "Deze documenten zijn niet stilzwijgend weggelaten; de bron kon de regels niet "
            "volledig als locatiegefilterde artikelen leveren."
        )
        for failure in failures:
            word.add_paragraph(failure, style="List Bullet")
    elif not articles:
        word.add_paragraph("De bron heeft geen afzonderlijke artikelen voor deze locatie teruggegeven.")

    output = BytesIO()
    word.save(output)
    return output.getvalue(), len(articles), len(failures)


class Handler(BaseHTTPRequestHandler):
    def send_json(self, status: int, payload: dict[str, Any]) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self) -> None:
        parsed = urllib.parse.urlsplit(self.path)
        if parsed.path == "/":
            body = HTML.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        if parsed.path == "/api/suggest":
            query = urllib.parse.parse_qs(parsed.query).get("q", [""])[0].strip()
            if len(query) < 3:
                self.send_json(400, {"error": "Vul minimaal drie tekens in."})
                return
            try:
                suggestions = get_suggestions(query)
            except UpstreamError as error:
                self.send_json(502, {"error": str(error)})
                return
            self.send_json(200, {"suggestions": suggestions})
            return
        self.send_json(404, {"error": "Niet gevonden."})

    def do_POST(self) -> None:
        path = urllib.parse.urlsplit(self.path).path
        if path not in {"/api/documents", "/api/word"}:
            self.send_json(404, {"error": "Niet gevonden."})
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            if length < 1 or length > 10_000:
                raise ValueError("Ongeldige aanvraaggrootte.")
            payload = json.loads(self.rfile.read(length))
            if not isinstance(payload, dict):
                raise ValueError("Ongeldige aanvraag.")
            address_id = payload.get("id")
            label = payload.get("label")
            if not isinstance(address_id, str) or not isinstance(label, str):
                raise ValueError("Selecteer eerst een adres uit de suggesties.")
        except (ValueError, json.JSONDecodeError) as error:
            self.send_json(400, {"error": str(error)})
            return
        if path == "/api/word":
            include_chapter_22 = payload.get("includeChapter22", False)
            if not isinstance(include_chapter_22, bool):
                self.send_json(400, {"error": "De keuze voor hoofdstuk 22 is ongeldig."})
                return
            include_draft_versions = payload.get("includeDraftVersions", False)
            if not isinstance(include_draft_versions, bool):
                self.send_json(400, {"error": "De keuze voor ontwerpversies is ongeldig."})
                return
            include_adopted_versions = payload.get("includeAdoptedVersions", False)
            if not isinstance(include_adopted_versions, bool):
                self.send_json(400, {"error": "De keuze voor vastgestelde versies is ongeldig."})
                return
            selected_document_ids_value = payload.get("selectedDocumentIds")
            if selected_document_ids_value is not None and (
                not isinstance(selected_document_ids_value, list)
                or not all(isinstance(item, str) for item in selected_document_ids_value)
            ):
                self.send_json(400, {"error": "De documentselectie is ongeldig."})
                return
            selected_document_ids = (
                set(selected_document_ids_value)
                if selected_document_ids_value is not None
                else None
            )
            try:
                content, article_count, failure_count = create_word_document(
                    address_id,
                    label,
                    include_chapter_22,
                    selected_document_ids,
                    include_draft_versions,
                    include_adopted_versions,
                )
            except UpstreamError as error:
                self.send_json(502, {"error": str(error)})
                return
            safe_address = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", label).strip(" .")
            filename = f"RPL_{safe_address or 'adres'}.docx"
            encoded_filename = urllib.parse.quote(filename, safe="")
            fallback_filename = re.sub(r"[^A-Za-z0-9._-]", "_", filename)
            self.send_response(200)
            self.send_header(
                "Content-Type",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            self.send_header("Content-Length", str(len(content)))
            self.send_header(
                "Content-Disposition",
                f'attachment; filename="{fallback_filename}"; filename*=UTF-8\'\'{encoded_filename}',
            )
            self.send_header("X-Download-Filename", encoded_filename)
            self.send_header("X-Article-Count", str(article_count))
            self.send_header("X-Incomplete-Documents", str(failure_count))
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            self.wfile.write(content)
            return
        try:
            result = fetch_documents(address_id, label)
        except UpstreamError as error:
            self.send_json(502, {"error": str(error)})
            return
        self.send_json(200, result)

    def log_message(self, format: str, *args: Any) -> None:
        with REQUEST_LOCK:
            super().log_message(format, *args)


def main() -> None:
    host = "127.0.0.1"
    port = 8000
    server = ThreadingHTTPServer((host, port), Handler)
    print(f"AdresRegels is beschikbaar op http://{host}:{port}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nServer wordt afgesloten.")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()